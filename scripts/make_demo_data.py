"""生成演示数据，用于快速验证环境与各功能是否正常。

用法::

    python scripts/make_demo_data.py            # 生成到 output/demo/inbox
    python scripts/make_demo_data.py --target output/demo/inbox
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import pandas as pd  # noqa: E402
from docx import Document  # noqa: E402

STAFF = [
    {"姓名": "张三", "联系电话": "13800000000", "家庭住址": "南京市玄武区某街道 1 号"},
    {"姓名": "李四", "联系电话": "13900000000", "家庭住址": "南京市鼓楼区某街道 2 号"},
]


def build(target: Path) -> list[Path]:
    target.mkdir(parents=True, exist_ok=True)
    (target / "子目录").mkdir(exist_ok=True)
    created: list[Path] = []

    # 1) 报表类：两份 Excel + 一份 CSV（列名略有差异，验证字段合并）
    frame_a = pd.DataFrame(STAFF)
    frame_b = pd.DataFrame([{"姓名": "王五", "电话": "13700000000", "住址": "苏州市姑苏区某街道 3 号"}])
    for name, frame in (("7月报表.xlsx", frame_a), ("8月报表.xlsx", frame_b)):
        path = target / name
        frame.to_excel(path, index=False)
        created.append(path)

    csv_path = target / "子目录" / "补充名单.csv"
    frame_a.to_csv(csv_path, index=False, encoding="utf-8-sig")
    created.append(csv_path)

    # 2) 文档类：Word 登记表、纯文本通知、Markdown 说明
    docx_path = target / "入户登记表.docx"
    document = Document()
    document.add_paragraph("姓名：赵六")
    document.add_paragraph("联系电话：13600000000")
    document.add_paragraph("家庭住址：无锡市梁溪区某街道 4 号")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "联系电话"
    table.cell(1, 0).text = "孙七"
    table.cell(1, 1).text = "13500000000"
    document.save(docx_path)
    created.append(docx_path)

    text_path = target / "工作通知.txt"
    text_path.write_text("各位同事：\n本周五下午三点在会议室集中学习新系统。\n", encoding="utf-8")
    created.append(text_path)

    md_path = target / "情况说明.md"
    md_path.write_text("# 情况说明\n\n本月累计办理业务 128 件，无异常。\n", encoding="utf-8")
    created.append(md_path)

    return created


def main() -> int:
    parser = argparse.ArgumentParser(description="生成演示数据")
    parser.add_argument("--target", default="output/demo/inbox", help="输出目录")
    args = parser.parse_args()

    target = Path(args.target)
    created = build(target)
    print(f"已生成 {len(created)} 份演示文件到：{target.resolve()}")
    for path in created:
        print(f"  - {path.relative_to(target)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
