"""多格式输入读取。

支持 Excel / CSV / TSV / Word / PDF / 纯文本 / Markdown / JSON / HTML，
并支持把整个目录（含子目录）批量读取为记录集合。
"""

from __future__ import annotations

import importlib.util
import json
import re
from collections.abc import Iterable, Iterator
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from docx import Document

from .formats import IMAGE_INPUT_SUFFIXES, INPUT_SUFFIXES, normalize_suffix

SOURCE_COLUMN = "源文件"
CONTENT_COLUMN = "内容"
PAGES_COLUMN = "页数"

EXCEL_SUFFIXES = frozenset({".xlsx", ".xlsm", ".xls"})
CSV_SUFFIXES = frozenset({".csv", ".tsv"})
TEXT_SUFFIXES = frozenset({".txt", ".md", ".log"})
HTML_SUFFIXES = frozenset({".html", ".htm"})

# OCR 引擎首次加载模型需要数秒，进程内复用实例
_OCR_ENGINE: Any = None

_ENCODINGS: tuple[str, ...] = ("utf-8-sig", "utf-8", "gb18030", "gbk", "big5", "latin-1")

# "字段：内容" 形式；字段名限制在 40 字符内，避免把整段正文误判为键值
_KEY_VALUE_PATTERN = re.compile(r"^\s*([^:：]{1,40}?)\s*[:：]\s*(.+?)\s*$")


# ------------------------------------------------------------------ 单文件读取
def read_excel(path: Path, sheet_name: int | str = 0) -> list[dict[str, Any]]:
    """读取 Excel 的第一个（或指定）工作表。"""
    try:
        frame = pd.read_excel(path, sheet_name=sheet_name, dtype=object)
    except ImportError as error:  # .xls 需要 xlrd
        raise ValueError(
            f"读取 {path.name} 失败：旧版 .xls 需要安装 xlrd（pip install xlrd）。原始错误：{error}"
        ) from error
    if isinstance(frame, dict):  # sheet_name=None 时返回字典
        frames = list(frame.values())
        frame = frames[0] if frames else pd.DataFrame()
    return _records_from_frame(frame)


def read_csv(path: Path, sep: str | None = None) -> list[dict[str, Any]]:
    """读取 CSV/TSV，自动尝试常见中文编码。"""
    separator = sep or ("\t" if normalize_suffix(path) == ".tsv" else ",")
    last_error: Exception | None = None
    for encoding in _ENCODINGS:
        try:
            frame = pd.read_csv(path, sep=separator, dtype=object, encoding=encoding)
        except UnicodeDecodeError as error:
            last_error = error
            continue
        except pd.errors.EmptyDataError:
            return []
        except pd.errors.ParserError as error:
            last_error = error
            continue
        return _records_from_frame(frame)
    raise ValueError(f"读取 {path.name} 失败，无法识别文件编码：{last_error}")


def read_word(path: Path) -> list[dict[str, Any]]:
    """读取 Word：段落中的 ``字段：内容`` 与文档内的表格都会被解析。"""
    document = Document(path)
    records: list[dict[str, Any]] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        pair = _split_key_value(text)
        if pair is not None:
            key, value = pair
            records.append({key: value})

    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if len(rows) < 2:
            continue
        headers = rows[0]
        for values in rows[1:]:
            records.append(dict(zip(headers, values, strict=False)))

    merged = _merge_records(records)
    if merged:
        return merged

    text = "\n".join(p.text.strip() for p in document.paragraphs if p.text.strip())
    return [{CONTENT_COLUMN: text}] if text else []


def read_pdf(
    path: Path,
    add_page_number: bool = False,
    *,
    use_ocr: bool = True,
) -> list[dict[str, Any]]:
    """读取 PDF：优先使用文本层；无文本层（扫描件）时改用 OCR。"""
    try:
        from pypdf import PdfReader
    except ImportError as error:  # pragma: no cover - 取决于运行环境
        raise ValueError("读取 PDF 需要安装 pypdf（pip install pypdf）。") from error

    reader = PdfReader(str(path))
    pages = [((page.extract_text() or "").strip()) for page in reader.pages]

    if add_page_number:
        records: list[dict[str, Any]] = []
        for index, text in enumerate(pages, start=1):
            if not text:
                continue
            for item in _records_from_text(text):
                records.append({SOURCE_COLUMN: path.name, "页码": index, **item})
        return records

    text = "\n".join(page for page in pages if page)
    if not text and use_ocr:
        text = ocr_pdf_text(path)  # 扫描件：逐页渲染后识别
    records = _records_from_text(text)
    if records:
        records[0][PAGES_COLUMN] = len(reader.pages)
    return records


def read_image(path: Path) -> list[dict[str, Any]]:
    """识别图片中的文字，再按文本规则解析为记录。"""
    try:
        from PIL import Image
    except ImportError as error:
        raise ValueError("读取图片需要安装 Pillow（pip install Pillow）。") from error

    with Image.open(path) as image:
        text = ocr_image_text(image)
    if not text:
        return []
    return _records_from_text(text)


def ocr_available() -> bool:
    """是否已安装 OCR 组件（rapidocr-onnxruntime），供界面与脚本提示使用。"""
    return importlib.util.find_spec("rapidocr_onnxruntime") is not None


def ocr_image_text(image: Any, *, engine: Any = None) -> str:
    """对单张图片执行 OCR，返回按行拼接的文本。"""
    recognizer = engine if engine is not None else _ocr_engine()
    result, _ = recognizer(np.asarray(image))
    if not result:
        return ""
    return "\n".join(str(item[1]) for item in result).strip()


def ocr_pdf_text(path: Path, *, scale: float = 2.0, max_pages: int | None = None) -> str:
    """把 PDF 页面渲染为位图后逐页识别（用于扫描件）。"""
    try:
        import pypdfium2 as pdfium
    except ImportError as error:
        raise ValueError(
            "识别扫描版 PDF 需要安装 pypdfium2（pip install pypdfium2）。"
        ) from error

    recognizer = _ocr_engine()
    document = pdfium.PdfDocument(str(path))
    total = len(document) if max_pages is None else min(len(document), max_pages)
    chunks: list[str] = []
    for index in range(total):
        bitmap = document[index].render(scale=scale)
        text = ocr_image_text(bitmap.to_pil(), engine=recognizer)
        if text:
            chunks.append(text)
    return "\n".join(chunks)


def _ocr_engine() -> Any:
    """惰性创建并复用 OCR 引擎。"""
    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
        except ImportError as error:
            raise ValueError(
                "识别图片或扫描件需要 OCR 组件："
                "pip install rapidocr-onnxruntime pypdfium2 Pillow"
            ) from error
        _OCR_ENGINE = RapidOCR()
    return _OCR_ENGINE


def read_text(path: Path) -> list[dict[str, Any]]:
    """读取纯文本 / Markdown：解析 ``字段：内容``，否则返回整篇内容。"""
    return _records_from_text(read_text_content(path))


def read_text_content(path: Path) -> str:
    """按常见编码读取文本文件的原始内容。"""
    last_error: Exception | None = None
    for encoding in _ENCODINGS:
        try:
            return path.read_text(encoding=encoding)
        except (UnicodeDecodeError, LookupError) as error:
            last_error = error
            continue
    raise ValueError(f"读取 {path.name} 失败，无法识别文件编码：{last_error}")


def read_json(path: Path) -> list[dict[str, Any]]:
    """读取 JSON：``[{...}]`` 与 ``{...}`` 均受支持。"""
    data = json.loads(read_text_content(path))
    if isinstance(data, list):
        if all(isinstance(item, dict) for item in data):
            return [dict(item) for item in data]
        return [{CONTENT_COLUMN: json.dumps(item, ensure_ascii=False)} for item in data]
    if isinstance(data, dict):
        if data and all(not isinstance(value, (dict, list)) for value in data.values()):
            return [dict(data)]
        return [{CONTENT_COLUMN: json.dumps(data, ensure_ascii=False, indent=2)}]
    return [{CONTENT_COLUMN: json.dumps(data, ensure_ascii=False)}]


def read_html(path: Path) -> list[dict[str, Any]]:
    """读取 HTML 中的第一张表格；没有表格时退化为整篇纯文本。"""
    try:
        tables = pd.read_html(str(path))
    except ImportError as error:  # pragma: no cover
        raise ValueError("读取 HTML 表格需要安装 lxml（pip install lxml）。") from error
    except ValueError:
        tables = []
    if tables:
        frame = tables[0]
        frame.columns = [str(column) for column in frame.columns]
        return _records_from_frame(frame)

    raw = path.read_text(encoding="utf-8", errors="ignore")
    text = re.sub(r"<[^>]+>", " ", raw)
    text = re.sub(r"\s+", " ", text).strip()
    return [{CONTENT_COLUMN: text}] if text else []


# ------------------------------------------------------------------ 统一入口
def read_any(path: str | Path, source: str | None = None) -> list[dict[str, Any]]:
    """按后缀自动选择读取器。``source`` 非空时给每条记录标注来源文件名。"""
    target = Path(path)
    suffix = normalize_suffix(target)
    if suffix in EXCEL_SUFFIXES:
        records = read_excel(target)
    elif suffix in CSV_SUFFIXES:
        records = read_csv(target)
    elif suffix == ".docx":
        records = read_word(target)
    elif suffix == ".pdf":
        records = read_pdf(target)
    elif suffix in TEXT_SUFFIXES:
        records = read_text(target)
    elif suffix == ".json":
        records = read_json(target)
    elif suffix in HTML_SUFFIXES:
        records = read_html(target)
    elif suffix in IMAGE_INPUT_SUFFIXES:
        records = read_image(target)
    else:
        raise ValueError(
            f"不支持的输入格式：{suffix or target.name}，"
            f"目前支持 {'、'.join(sorted(INPUT_SUFFIXES))}"
        )
    if source:
        return [{SOURCE_COLUMN: source, **record} for record in records]
    return records


def read_input(path: Path) -> list[dict[str, Any]]:
    """向后兼容旧版 API。"""
    return read_any(path)


# ------------------------------------------------------------------ 目录批量
def iter_input_files(
    directory: str | Path,
    *,
    recursive: bool = True,
    patterns: Iterable[str] | None = None,
    include_temp: bool = False,
) -> list[Path]:
    """列出目录下所有可读取的文档，按路径排序。"""
    root = Path(directory)
    if not root.is_dir():
        raise ValueError(f"不是有效目录：{root}")

    matched: set[Path] = set()
    if patterns:
        for pattern in patterns:
            matched.update(root.rglob(pattern) if recursive else root.glob(pattern))
    else:
        matched.update(root.rglob("*") if recursive else root.glob("*"))

    files: list[Path] = []
    for path in sorted(matched):
        if not path.is_file():
            continue
        if normalize_suffix(path) not in INPUT_SUFFIXES:
            continue
        if not include_temp and (path.name.startswith("~$") or path.name.startswith(".")):
            continue
        files.append(path)
    return files


def read_folder(
    directory: str | Path,
    *,
    recursive: bool = True,
    patterns: Iterable[str] | None = None,
    source_column: str | None = SOURCE_COLUMN,
    relative: bool = True,
    skip_errors: bool = True,
    failed: list[tuple[Path, str]] | None = None,
) -> list[dict[str, Any]]:
    """批量读取目录下所有文档，合并为一个记录集合。

    默认跳过无法解析的文件（例如损坏文件），并把失败原因记录到 ``failed``，
    避免个别文件出错导致整批任务中断。
    """
    root = Path(directory)
    records: list[dict[str, Any]] = []
    for path in iter_input_files(root, recursive=recursive, patterns=patterns):
        label = str(path.relative_to(root)) if relative else path.name
        try:
            items = read_any(path)
        except Exception as error:  # noqa: BLE001 - 单份文件失败不应中断整批
            if not skip_errors:
                raise
            if failed is not None:
                failed.append((path, str(error)))
            continue
        if source_column is None:
            records.extend(items)
        else:
            records.extend({source_column: label, **item} for item in items)
    return records


# ------------------------------------------------------------------ 内部工具
def _records_from_frame(frame: pd.DataFrame) -> list[dict[str, Any]]:
    frame = frame.where(frame.notna(), "")
    return [
        {str(key).strip(): value for key, value in record.items()}
        for record in frame.to_dict(orient="records")
    ]


def _records_from_text(text: str) -> list[dict[str, Any]]:
    """把正文解析为键值记录；解析不到键值时返回整篇内容。"""
    cleaned = text.strip()
    if not cleaned:
        return []

    pairs: dict[str, Any] = {}
    unstructured: list[str] = []
    for line in cleaned.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        pair = _split_key_value(stripped)
        if pair is None:
            unstructured.append(stripped)
        else:
            key, value = pair
            pairs[key] = value

    if pairs and not unstructured:
        return [pairs]
    if pairs:
        return [{**pairs, CONTENT_COLUMN: "\n".join(unstructured)}]
    return [{CONTENT_COLUMN: cleaned}]


def _split_key_value(text: str) -> tuple[str, str] | None:
    """识别 ``字段：内容`` 形式的文本行，识别失败返回 ``None``。"""
    match = _KEY_VALUE_PATTERN.match(text)
    if match:
        return match.group(1).strip(), match.group(2).strip()

    # 退化分支：全角空格或连续两个空格分隔的 "键 值"
    for separator in ("　", "  "):
        if separator in text:
            key, _, value = text.partition(separator)
            key, value = key.strip(), value.strip()
            if key and value and len(key) <= 20 and " " not in key:
                return key, value
    return None


def _merge_records(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """把键值段落合并为一条记录；表格行保持各自记录。"""
    if not records:
        return []
    merged: dict[str, Any] = {}
    for record in records:
        if len(record) == 1:
            merged.update(record)
        else:
            return records
    return [merged]


def is_plain_content(records: list[dict[str, Any]]) -> bool:
    """判断记录是否只是"整篇正文"（仅含内容/页数字段）。"""
    return bool(records) and all(
        set(record) <= {CONTENT_COLUMN, PAGES_COLUMN} for record in records
    )


def content_of(records: list[dict[str, Any]]) -> str:
    """取出记录中的正文内容。"""
    return "\n\n".join(str(record.get(CONTENT_COLUMN, "")) for record in records).strip()


def iter_supported_suffixes() -> Iterator[str]:
    return iter(sorted(INPUT_SUFFIXES))
