"""多格式输出。

支持把记录集合写成 Excel / CSV / Word / 纯文本 / Markdown / JSON / HTML，
并支持对已有 Excel / CSV 追加写入（大批量分批处理时使用）。
"""

from __future__ import annotations

import json
from datetime import datetime
from html import escape
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document

from .formats import OUTPUT_SUFFIXES, normalize_suffix

EXCEL_SUFFIXES = frozenset({".xlsx", ".xlsm"})


def write_output(
    records: list[dict[str, Any]],
    headers: list[str],
    output_path: str | Path,
    *,
    title: str | None = None,
    sheet_name: str = "Sheet1",
    append: bool = False,
) -> Path:
    """把记录写入指定格式的文件，返回最终路径。"""
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    suffix = normalize_suffix(target)
    normalized = [dict(record) for record in records]

    if suffix in EXCEL_SUFFIXES:
        _write_excel(normalized, headers, target, sheet_name=sheet_name, append=append)
    elif suffix == ".csv":
        _write_csv(normalized, headers, target, append=append)
    elif suffix == ".docx":
        _write_docx(normalized, headers, target, title=title)
    elif suffix == ".md":
        _write_markdown(normalized, headers, target, title=title)
    elif suffix == ".txt":
        _write_text(normalized, headers, target, title=title)
    elif suffix == ".json":
        _write_json(normalized, headers, target)
    elif suffix in {".html", ".htm"}:
        _write_html(normalized, headers, target, title=title)
    else:
        raise ValueError(
            f"不支持的输出格式：{suffix or target.name}，"
            f"请使用 {'、'.join(sorted(OUTPUT_SUFFIXES))}"
        )
    return target


def write_text_document(
    text: str,
    output_path: str | Path,
    *,
    title: str | None = None,
) -> Path:
    """把整篇正文写入 txt / md / docx / html，用于"原文直转"。"""
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    suffix = normalize_suffix(target)
    body = text or ""

    if suffix in {".txt", ".log"}:
        target.write_text(body, encoding="utf-8")
    elif suffix == ".md":
        parts = [f"# {title}", "", body] if title else [body]
        target.write_text("\n".join(parts).strip() + "\n", encoding="utf-8")
    elif suffix == ".docx":
        document = Document()
        if title:
            document.add_heading(title, level=0)
        for line in body.splitlines():
            document.add_paragraph(line)
        document.save(target)
    elif suffix in {".html", ".htm"}:
        heading = f"<h1>{escape(title)}</h1>" if title else ""
        target.write_text(_HTML_TEMPLATE.format(title=escape(title or ""), body=heading + escape(body)), encoding="utf-8")
    elif suffix == ".json":
        target.write_text(
            json.dumps({"标题": title, "内容": body}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    else:
        # Excel/CSV 无法存放纯文本，按行拆成单字段记录
        records = [{"内容": line} for line in body.splitlines()]
        write_output(records, ["内容"], target, title=title)
    return target


# ---------------------------------------------------------------- 各格式实现
def _write_excel(
    records: list[dict[str, Any]],
    headers: list[str],
    target: Path,
    *,
    sheet_name: str,
    append: bool,
) -> None:
    if append and target.exists():
        from openpyxl import load_workbook

        workbook = load_workbook(target)
        sheet = (
            workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.create_sheet(sheet_name)
        )
        if sheet.max_row == 1 and all(cell.value in (None, "") for cell in sheet[1]):
            sheet.delete_rows(1)
        if sheet.max_row == 0:
            sheet.append(list(headers))
        for record in records:
            sheet.append([record.get(header, "") for header in headers])
        workbook.save(target)
        return

    frame = pd.DataFrame(records, columns=headers)
    frame.to_excel(target, index=False, sheet_name=sheet_name)


def _write_csv(
    records: list[dict[str, Any]],
    headers: list[str],
    target: Path,
    *,
    append: bool,
) -> None:
    frame = pd.DataFrame(records, columns=headers)
    mode = "a" if append and target.exists() else "w"
    frame.to_csv(
        target,
        index=False,
        encoding="utf-8-sig",  # BOM 让 Excel 正确识别中文
        mode=mode,
        header=not (append and target.exists() and target.stat().st_size > 0),
    )


def _write_docx(
    records: list[dict[str, Any]],
    headers: list[str],
    target: Path,
    *,
    title: str | None,
) -> None:
    document = Document()
    if title:
        document.add_heading(title, level=0)

    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers, strict=False):
        cell.text = str(header)

    for record in records:
        cells = table.add_row().cells
        for cell, header in zip(cells, headers, strict=False):
            cell.text = str(record.get(header, ""))

    document.add_paragraph(f"共 {len(records)} 条记录 · 生成时间 {_now()}")
    document.save(target)


def _write_markdown(
    records: list[dict[str, Any]],
    headers: list[str],
    target: Path,
    *,
    title: str | None,
) -> None:
    lines: list[str] = []
    if title:
        lines.extend([f"# {title}", ""])
    if headers:
        lines.append("| " + " | ".join(_markdown_cell(header) for header in headers) + " |")
        lines.append("| " + " | ".join("---" for _ in headers) + " |")
        for record in records:
            cells = [_markdown_cell(record.get(header, "")) for header in headers]
            lines.append("| " + " | ".join(cells) + " |")
    lines.append("")
    lines.append(f"> 共 {len(records)} 条记录 · 生成时间 {_now()}")
    target.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_text(
    records: list[dict[str, Any]],
    headers: list[str],
    target: Path,
    *,
    title: str | None,
) -> None:
    lines: list[str] = []
    if title:
        lines.extend([title, "=" * _display_width(title), ""])
    rows = [[str(record.get(header, "")) for header in headers] for record in records]
    widths = [
        max([_display_width(str(header))] + [_display_width(row[i]) for row in rows]) if rows else _display_width(str(header))
        for i, header in enumerate(headers)
    ]
    lines.append("  ".join(_pad(str(header), widths[i]) for i, header in enumerate(headers)))
    lines.append("  ".join("-" * width for width in widths))
    for row in rows:
        lines.append("  ".join(_pad(value, widths[i]) for i, value in enumerate(row)))
    lines.extend(["", f"共 {len(records)} 条记录 · 生成时间 {_now()}"])
    target.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_json(records: list[dict[str, Any]], headers: list[str], target: Path) -> None:
    payload = {
        "生成时间": _now(),
        "记录数": len(records),
        "表头": list(headers),
        "记录": [{header: record.get(header, "") for header in headers} for record in records],
    }
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_html(
    records: list[dict[str, Any]],
    headers: list[str],
    target: Path,
    *,
    title: str | None,
) -> None:
    heading = f"<h1>{escape(title)}</h1>" if title else ""
    head_cells = "".join(f"<th>{escape(str(header))}</th>" for header in headers)
    body_rows = []
    for record in records:
        cells = "".join(
            f"<td>{escape(str(record.get(header, '')))}</td>" for header in headers
        )
        body_rows.append(f"<tr>{cells}</tr>")
    table = (
        f"<table><thead><tr>{head_cells}</tr></thead><tbody>{''.join(body_rows)}</tbody></table>"
    )
    footer = f"<p class=\"meta\">共 {len(records)} 条记录 · 生成时间 {_now()}</p>"
    target.write_text(
        _HTML_TEMPLATE.format(title=escape(title or "导出结果"), body=heading + table + footer),
        encoding="utf-8",
    )


# ---------------------------------------------------------------- 小工具
_HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>
  body {{ font-family: "Microsoft YaHei", "PingFang SC", Arial, sans-serif; margin: 24px; color: #222; }}
  table {{ border-collapse: collapse; width: 100%; }}
  th, td {{ border: 1px solid #bbb; padding: 6px 10px; text-align: left; font-size: 14px; }}
  th {{ background: #f2f4f7; }}
  tr:nth-child(even) td {{ background: #fafbfc; }}
  .meta {{ color: #777; font-size: 12px; margin-top: 12px; }}
</style>
</head>
<body>
{body}
</body>
</html>
"""


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _display_width(text: str) -> int:
    """估算显示宽度，中文等全角字符按 2 计算，用于文本表对齐。"""
    import unicodedata

    return sum(
        2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1 for char in str(text)
    )


def _pad(text: str, width: int) -> str:
    padding = max(width - _display_width(text), 0)
    return str(text) + " " * padding


def _markdown_cell(value: Any) -> str:
    return str(value).replace("|", "\\|").replace("\n", "<br>")
