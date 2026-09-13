"""测试共用的样例数据。"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from docx import Document

PEOPLE = [
    {"姓名": "张三", "联系电话": "13800000000", "家庭住址": "南京市玄武区"},
    {"姓名": "李四", "联系电话": "13900000000", "家庭住址": "南京市鼓楼区"},
]


@pytest.fixture
def people_records() -> list[dict[str, str]]:
    return [dict(item) for item in PEOPLE]


@pytest.fixture
def xlsx_file(tmp_path: Path, people_records: list[dict[str, str]]) -> Path:
    path = tmp_path / "报表A.xlsx"
    pd.DataFrame(people_records).to_excel(path, index=False)
    return path


@pytest.fixture
def xlsx_file_b(tmp_path: Path) -> Path:
    path = tmp_path / "报表B.xlsx"
    pd.DataFrame(
        [
            {"姓名": "王五", "联系电话": "13700000000", "家庭住址": "苏州市姑苏区"},
        ]
    ).to_excel(path, index=False)
    return path


@pytest.fixture
def csv_file(tmp_path: Path, people_records: list[dict[str, str]]) -> Path:
    path = tmp_path / "名单.csv"
    pd.DataFrame(people_records).to_csv(path, index=False, encoding="utf-8-sig")
    return path


@pytest.fixture
def gbk_csv_file(tmp_path: Path) -> Path:
    path = tmp_path / "旧系统导出.csv"
    pd.DataFrame([{"姓名": "赵六", "电话": "13600000000"}]).to_csv(
        path, index=False, encoding="gbk"
    )
    return path


@pytest.fixture
def docx_file(tmp_path: Path) -> Path:
    path = tmp_path / "登记表.docx"
    document = Document()
    document.add_paragraph("姓名：王五")
    document.add_paragraph("联系电话：13700000000")
    document.add_paragraph("家庭住址：苏州市姑苏区")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "联系电话"
    table.cell(1, 0).text = "孙七"
    table.cell(1, 1).text = "13500000000"
    document.save(path)
    return path


@pytest.fixture
def text_file(tmp_path: Path) -> Path:
    path = tmp_path / "说明.txt"
    path.write_text("姓名：周八\n联系电话：13400000000\n", encoding="utf-8")
    return path


@pytest.fixture
def plain_text_file(tmp_path: Path) -> Path:
    path = tmp_path / "通知.txt"
    path.write_text("各位同事：\n本周五下午三点在会议室集中学习。\n", encoding="utf-8")
    return path


@pytest.fixture
def json_file(tmp_path: Path) -> Path:
    path = tmp_path / "数据.json"
    path.write_text(
        '[{"姓名": "吴九", "联系电话": "13300000000"}]', encoding="utf-8"
    )
    return path


@pytest.fixture
def inbox(tmp_path: Path, people_records: list[dict[str, str]]) -> Path:
    """模拟真实收件目录：多种格式混放，含子目录。"""
    root = tmp_path / "inbox"
    (root / "子目录").mkdir(parents=True)

    pd.DataFrame(people_records).to_excel(root / "报表A.xlsx", index=False)
    pd.DataFrame([{"姓名": "王五", "电话": "13700000000"}]).to_csv(
        root / "子目录" / "报表B.csv", index=False, encoding="utf-8-sig"
    )

    document = Document()
    document.add_paragraph("姓名：王五")
    document.add_paragraph("联系电话：13700000000")
    document.save(root / "登记表.docx")

    (root / "说明.txt").write_text("姓名：周八\n联系电话：13400000000\n", encoding="utf-8")
    (root / "图片说明.md").write_text("# 说明\n\n这是一份无结构正文。\n", encoding="utf-8")
    (root / "临时文件.docx").write_bytes(b"not a real docx")  # 用于验证失败不中断整批
    return root
