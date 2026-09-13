"""格式转换测试。"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from docx import Document

from information_mapper.converters import (
    convert_file,
    convert_folder,
    normalize_target_format,
    summarize,
)


def test_normalize_target_format() -> None:
    assert normalize_target_format("csv") == ".csv"
    assert normalize_target_format(".XLSX") == ".xlsx"
    with pytest.raises(ValueError, match="不支持的目标格式"):
        normalize_target_format("pptx")


def test_convert_xlsx_to_csv(xlsx_file: Path, tmp_path: Path) -> None:
    result = convert_file(xlsx_file, tmp_path / "结果.csv")
    assert result.ok
    frame = pd.read_csv(result.target, encoding="utf-8-sig")
    assert frame.iloc[0]["姓名"] == "张三"


def test_convert_csv_to_docx(csv_file: Path, tmp_path: Path) -> None:
    result = convert_file(csv_file, tmp_path / "结果.docx")
    assert result.ok
    document = Document(str(result.target))
    assert document.tables[0].rows[1].cells[0].text == "张三"


def test_convert_plain_text_keeps_original_body(plain_text_file: Path, tmp_path: Path) -> None:
    result = convert_file(plain_text_file, tmp_path / "通知.docx")
    assert result.ok
    document = Document(str(result.target))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "会议室" in text


def test_convert_with_mapping(xlsx_file: Path, tmp_path: Path) -> None:
    mapping = {
        "target_headers": ["联系电话", "姓名"],
        "field_mapping": {"联系电话": ["电话", "联系电话"]},
    }
    result = convert_file(xlsx_file, tmp_path / "映射.csv", mapping=mapping)
    assert result.ok
    frame = pd.read_csv(result.target, encoding="utf-8-sig")
    assert list(frame.columns) == ["联系电话", "姓名"]


def test_convert_folder_batch_reports_failures(inbox: Path, tmp_path: Path) -> None:
    results = convert_folder(inbox, tmp_path / "out", "csv")
    assert any(item.ok for item in results)
    assert any(not item.ok for item in results)
    assert "失败" in summarize(results)


def test_convert_folder_keeps_relative_structure(inbox: Path, tmp_path: Path) -> None:
    results = convert_folder(inbox, tmp_path / "out", "csv")
    targets = {item.target for item in results if item.ok}
    assert tmp_path / "out" / "报表A.csv" in targets
    assert tmp_path / "out" / "子目录" / "报表B.csv" in targets


def test_convert_folder_avoids_overwriting_existing_files(inbox: Path, tmp_path: Path) -> None:
    outcome = tmp_path / "out"
    convert_folder(inbox, outcome, "csv")
    second = convert_folder(inbox, outcome, "csv")
    assert any("(1)" in item.target.name for item in second if item.ok)


def test_convert_folder_can_filter_patterns(inbox: Path, tmp_path: Path) -> None:
    results = convert_folder(inbox, tmp_path / "out", "md", patterns=["*.csv"])
    assert [item.source.name for item in results] == ["报表B.csv"]


def test_convert_missing_input_raises(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="找不到输入文件"):
        convert_file(tmp_path / "不存在.xlsx", tmp_path / "结果.csv")


def test_convert_rejects_unknown_target(tmp_path: Path, xlsx_file: Path) -> None:
    with pytest.raises(ValueError, match="不支持的输出格式"):
        convert_file(xlsx_file, tmp_path / "结果.ppt")
