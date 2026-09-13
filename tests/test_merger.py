"""批量合并测试。"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from docx import Document

from information_mapper.formats import normalize_suffix
from information_mapper.merger import (
    collect_headers,
    merge_any,
    merge_tables,
    merge_text_documents,
    merge_word_documents,
    pick_mode,
    resolve_inputs,
)
from information_mapper.readers import iter_input_files


def test_merge_tables_adds_source_column(
    xlsx_file: Path, xlsx_file_b: Path, csv_file: Path, tmp_path: Path
) -> None:
    result = merge_tables([xlsx_file, xlsx_file_b, csv_file], tmp_path / "汇总.xlsx")
    assert result.record_count == 5
    frame = pd.read_excel(result.output)
    assert list(frame.columns)[0] == "源文件"
    assert set(frame["源文件"]) == {"报表A.xlsx", "报表B.xlsx", "名单.csv"}


def test_merge_tables_can_disable_source_column(xlsx_file: Path, tmp_path: Path) -> None:
    result = merge_tables([xlsx_file], tmp_path / "汇总.xlsx", source_column=None)
    frame = pd.read_excel(result.output)
    assert "源文件" not in frame.columns


def test_merge_tables_respects_header_order(xlsx_file: Path, tmp_path: Path) -> None:
    result = merge_tables([xlsx_file], tmp_path / "汇总.xlsx", headers=["联系电话", "姓名"])
    assert result.headers[:2] == ["联系电话", "姓名"]


def test_merge_tables_skips_broken_files(inbox: Path, tmp_path: Path) -> None:
    result = merge_tables(iter_input_files(inbox), tmp_path / "汇总.xlsx", headers=["姓名"])
    assert [path.name for path, _ in result.failed] == ["临时文件.docx"]
    assert result.headers[0] == "姓名"  # 指定的表头排在最前，其余字段自动补在后面
    assert "源文件" in result.headers
    assert result.record_count > 0


def test_merge_any_accepts_directory(inbox: Path, tmp_path: Path) -> None:
    result = merge_any([inbox], tmp_path / "汇总.xlsx")
    assert normalize_suffix(result.output) == ".xlsx"
    assert result.record_count > 0
    assert result.output.exists()


def test_merge_any_can_write_markdown(inbox: Path, tmp_path: Path) -> None:
    result = merge_any([inbox], tmp_path / "汇总.md", patterns=["*.csv"])
    assert "报表B.csv" in result.output.read_text(encoding="utf-8")


def test_merge_any_word_mode(tmp_path: Path) -> None:
    first = tmp_path / "第一份.docx"
    second = tmp_path / "第二份.docx"
    for path, text in ((first, "第一份正文"), (second, "第二份正文")):
        document = Document()
        document.add_paragraph(text)
        document.save(path)

    result = merge_any([first, second], tmp_path / "合集.docx", title="合集")
    document = Document(str(result.output))
    body = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "第一份正文" in body
    assert "第二份正文" in body
    assert "合集" in body


def test_merge_word_documents_preserves_tables(docx_file: Path, tmp_path: Path) -> None:
    result = merge_word_documents([docx_file], tmp_path / "合集.docx")
    document = Document(str(result.output))
    assert len(document.tables) == 1
    assert document.tables[0].rows[1].cells[0].text == "孙七"


def test_merge_text_documents(tmp_path: Path, plain_text_file: Path, text_file: Path) -> None:
    result = merge_text_documents(
        [plain_text_file, text_file], tmp_path / "合集.md", title="材料合集"
    )
    content = result.output.read_text(encoding="utf-8")
    assert "会议室" in content
    assert "周八" in content
    assert "## 通知" in content


def test_merge_empty_input_raises(tmp_path: Path) -> None:
    (tmp_path / "空目录").mkdir()
    with pytest.raises(ValueError, match="没有找到可合并的输入文档"):
        merge_any([tmp_path / "空目录"], tmp_path / "汇总.xlsx")


def test_merge_missing_input_raises(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="找不到输入"):
        merge_any([tmp_path / "不存在"], tmp_path / "汇总.xlsx")


def test_resolve_inputs_deduplicates(inbox: Path) -> None:
    files = resolve_inputs([inbox, inbox / "报表A.xlsx"])
    names = [path.name for path in files]
    assert names.count("报表A.xlsx") == 1


def test_collect_headers_prefers_given_order() -> None:
    records = [{"姓名": "张三", "电话": "1"}, {"地址": "南京"}]
    assert collect_headers(records, ["地址"]) == ["地址", "姓名", "电话"]


@pytest.mark.parametrize(
    ("output", "inputs", "expected"),
    [
        ("汇总.xlsx", ["a.csv"], "table"),
        ("汇总.docx", ["a.docx"], "word"),
        ("合集.md", ["a.txt", "b.md"], "text"),
        ("汇总.docx", ["a.xlsx"], "table"),
        ("汇总.docx", ["a.pdf", "b.txt"], "text"),
    ],
)
def test_pick_mode_auto(output: str, inputs: list[str], expected: str) -> None:
    paths = [Path(name) for name in inputs]
    assert pick_mode(paths, output, "auto") == expected
