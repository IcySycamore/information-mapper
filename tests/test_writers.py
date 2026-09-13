"""多格式输出测试。"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest

from information_mapper.writers import write_output, write_text_document

HEADERS = ["姓名", "联系电话"]
RECORDS = [{"姓名": "张三", "联系电话": "13800000000"}]


def test_write_xlsx_round_trip(tmp_path: Path) -> None:
    target = write_output(RECORDS, HEADERS, tmp_path / "结果.xlsx")
    frame = pd.read_excel(target, dtype=object)
    assert list(frame.columns) == HEADERS
    assert frame.iloc[0]["姓名"] == "张三"


def test_write_csv_uses_utf8_bom(tmp_path: Path) -> None:
    target = write_output(RECORDS, HEADERS, tmp_path / "结果.csv")
    assert target.read_bytes().startswith(b"\xef\xbb\xbf")
    frame = pd.read_csv(target, encoding="utf-8-sig", dtype=str)  # 避免电话号码被当成数字
    assert frame.iloc[0]["联系电话"] == "13800000000"


def test_write_csv_append_keeps_single_header(tmp_path: Path) -> None:
    target = tmp_path / "追加.csv"
    write_output(RECORDS, HEADERS, target)
    write_output([{"姓名": "李四", "联系电话": "13900000000"}], HEADERS, target, append=True)
    frame = pd.read_csv(target, encoding="utf-8-sig")
    assert len(frame) == 2
    assert list(frame.columns) == HEADERS


def test_write_docx_contains_table(tmp_path: Path) -> None:
    from docx import Document

    target = write_output(RECORDS, HEADERS, tmp_path / "结果.docx", title="汇总表")
    document = Document(str(target))
    assert document.paragraphs[0].text == "汇总表"
    table = document.tables[0]
    assert [cell.text for cell in table.rows[0].cells] == HEADERS
    assert table.rows[1].cells[0].text == "张三"


def test_write_markdown(tmp_path: Path) -> None:
    target = write_output(RECORDS, HEADERS, tmp_path / "结果.md")
    content = target.read_text(encoding="utf-8")
    assert "| 姓名 | 联系电话 |" in content
    assert "| 张三 | 13800000000 |" in content


def test_write_json_payload(tmp_path: Path) -> None:
    import json

    target = write_output(RECORDS, HEADERS, tmp_path / "结果.json")
    payload = json.loads(target.read_text(encoding="utf-8"))
    assert payload["记录数"] == 1
    assert payload["表头"] == HEADERS


def test_write_html_outputs_table(tmp_path: Path) -> None:
    target = write_output(RECORDS, HEADERS, tmp_path / "结果.html")
    content = target.read_text(encoding="utf-8")
    assert "<table>" in content
    assert "张三" in content


def test_write_text_aligns_columns(tmp_path: Path) -> None:
    target = write_output(RECORDS, HEADERS, tmp_path / "结果.txt")
    content = target.read_text(encoding="utf-8")
    assert "姓名" in content and "张三" in content
    assert "共 1 条记录" in content


def test_write_output_creates_missing_directories(tmp_path: Path) -> None:
    target = write_output(RECORDS, HEADERS, tmp_path / "深" / "层" / "结果.csv")
    assert target.exists()


def test_write_output_rejects_unknown_suffix(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="不支持的输出格式"):
        write_output(RECORDS, HEADERS, tmp_path / "结果.pdf")


def test_write_text_document_to_docx(tmp_path: Path) -> None:
    from docx import Document

    target = write_text_document("第一行\n第二行", tmp_path / "正文.docx", title="通知")
    document = Document(str(target))
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "通知" in texts[0]
    assert "第二行" in texts


def test_write_text_document_to_txt(tmp_path: Path) -> None:
    target = write_text_document("原文内容", tmp_path / "正文.txt")
    assert target.read_text(encoding="utf-8") == "原文内容"
